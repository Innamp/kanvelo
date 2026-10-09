import openpyxl
import json
import re

wb = openpyxl.load_workbook('Sprint Backlog Tracker - EI.xlsx', read_only=True, data_only=True)

def clean_val(v):
    if v is None:
        return ''
    s = str(v).strip()
    if s == 'None' or s == '#VALUE!' or s == '#REF!':
        return ''
    return s

def map_stage(status, update_val=''):
    comb = (status + ' ' + update_val).lower()
    if any(k in comb for k in ['closed', 'done', 'deployed', 'complete', 'finish']):
        return 'completed', 100
    if any(k in comb for k in ['review', 'testing', 'qa', 'stag', 'pvt', 'verification']):
        return 'review', 80
    if any(k in comb for k in ['develop', 'progress', 'development', 'sharing', 'on dev']):
        return 'in-progress', 45
    return 'todo', 0

def map_prio(row):
    # check rank prio or text
    prio_raw = clean_val(row[12] if len(row) > 12 else '')
    mandays = clean_val(row[26] if len(row) > 26 else '')
    
    if 'p0' in prio_raw.lower() or 'high' in prio_raw.lower():
        return 'High'
    if 'p2' in prio_raw.lower() or 'low' in prio_raw.lower():
        return 'Low'
    try:
        val = float(mandays)
        if val >= 20: return 'High'
        if val <= 5: return 'Low'
    except:
        pass
    return 'Medium'

sprints_data = {}

for sheet_name, sprint_id, sprint_title in [
    ('Sprint 10 Backlog with Prio', 'sprint-10', 'Sprint 10 (Current)'),
    ('Sprint 9 Backlog with Prio', 'sprint-9', 'Sprint 9'),
    ('Sprint 8  with Prio', 'sprint-8', 'Sprint 8'),
]:
    if sheet_name not in wb.sheetnames:
        continue
    ws = wb[sheet_name]
    tasks = []
    
    for i, row in enumerate(ws.iter_rows(values_only=True)):
        if i < 8:
            continue
        if not row or not any(row):
            continue
            
        objective = clean_val(row[5] if len(row) > 5 else '')
        background = clean_val(row[4] if len(row) > 4 else '')
        description = clean_val(row[6] if len(row) > 6 else '')
        
        title = objective or background or description
        if not title or len(title) < 3:
            continue
            
        platform = clean_val(row[1] if len(row) > 1 else '') or clean_val(row[2] if len(row) > 2 else '') or 'General'
        pm = clean_val(row[7] if len(row) > 7 else '') or clean_val(row[8] if len(row) > 8 else '') or 'Unassigned'
        status_raw = clean_val(row[15] if len(row) > 15 else '')
        update_raw = clean_val(row[19] if len(row) > 19 else (row[18] if len(row) > 18 else ''))
        jira_val = clean_val(row[25] if len(row) > 25 else '')
        mandays_val = clean_val(row[26] if len(row) > 26 else '')
        dependency = clean_val(row[3] if len(row) > 3 else '')
        
        stage, default_progress = map_stage(status_raw, update_raw)
        prio = map_prio(row)
        
        # generate task ID
        task_id = f"{sprint_id}-t{len(tasks)+1}"
        
        tasks.append({
            "id": task_id,
            "title": title[:100],
            "description": description or background or objective,
            "background": background,
            "objective": objective,
            "platform": platform,
            "assignee": pm,
            "priority": prio,
            "stage": stage, # 'todo', 'in-progress', 'review', 'completed'
            "progress": default_progress if default_progress > 0 else (100 if stage == 'completed' else (50 if stage == 'in-progress' else 0)),
            "mandays": mandays_val or "5",
            "jira": jira_val,
            "dependency": dependency,
            "statusRaw": status_raw,
            "commentsCount": (len(description) % 5) + 1,
            "attachmentsCount": (len(title) % 4)
        })
        
    sprints_data[sprint_id] = {
        "id": sprint_id,
        "title": sprint_title,
        "tasksCount": len(tasks),
        "tasks": tasks
    }

with open('tasks_data.json', 'w', encoding='utf-8') as f:
    json.dump(sprints_data, f, indent=2, ensure_ascii=False)

print(f"Successfully exported {len(sprints_data)} sprints to tasks_data.json")
for sid, sdata in sprints_data.items():
    print(f"- {sid}: {sdata['tasksCount']} tasks")
